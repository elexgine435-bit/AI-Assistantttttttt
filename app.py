"""ATS Resume Checker - Streamlit + Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-2.5-flash"
MAX_CHARS = 20000  # safety cap on text sent to the model

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. Evaluate the resume text strictly and honestly. \
Do not inflate scores. Only reference content that is actually in the resume. \
Return ONLY valid JSON matching the requested schema."""

SCHEMA_HINT = """Return JSON with exactly this structure:
{
  "overall_score": <integer 0-100>,
  "section_scores": {
    "formatting_and_structure": <0-100>,
    "keywords_and_skills": <0-100>,
    "work_experience_impact": <0-100>,
    "education_and_certifications": <0-100>,
    "readability_and_clarity": <0-100>
  },
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", ...],
  "missing_keywords": ["<string>", ...],
  "improvements": [
    {"section": "<resume section>", "priority": "High|Medium|Low",
     "issue": "<what is wrong>", "suggestion": "<specific fix>",
     "example": "<optional rewritten example line, or empty string>"}
  ]
}
Give 5-10 improvements, ordered by priority."""


# ---------------------------------------------------------------- extraction
def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")


# ------------------------------------------------------------ rule-based checks
SECTION_PATTERNS = {
    "Summary / Objective": r"\b(summary|objective|profile)\b",
    "Experience": r"\b(experience|employment|work history)\b",
    "Education": r"\beducation\b",
    "Skills": r"\bskills?\b",
    "Projects": r"\bprojects?\b",
}


def basic_checks(text: str) -> list:
    """Quick deterministic checks. Returns a list of (passed, message)."""
    words = len(text.split())
    checks = [
        (bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)), "Email address found"),
        (bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)), "Phone number found"),
        (bool(re.search(r"linkedin\.com|github\.com", text, re.I)),
         "LinkedIn / GitHub link found"),
        (bool(re.search(r"\d+\s?%|[$€£]\s?\d|\b\d+[kKmMxX]\+?\b|\b\d{1,3}(,\d{3})+\b", text)),
         "Quantified achievements (numbers, %, $) found"),
        (250 <= words <= 1200, f"Length is reasonable ({words} words; ideal 250-1200)"),
    ]
    for label, pattern in SECTION_PATTERNS.items():
        checks.append((bool(re.search(pattern, text, re.I)), f"'{label}' section detected"))
    return checks


# ------------------------------------------------------------------ AI analysis
def parse_json(raw: str) -> dict:
    """Parse model output as JSON, tolerating markdown fences / extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.S)
        if match:
            return json.loads(match.group(0))
        raise


def clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize(result: dict) -> dict:
    """Make sure the result has every field the UI needs."""
    sections = result.get("section_scores") or {}
    return {
        "overall_score": clamp(result.get("overall_score")),
        "section_scores": {str(k): clamp(v) for k, v in sections.items()},
        "summary": str(result.get("summary", "")),
        "strengths": [str(s) for s in result.get("strengths", []) or []],
        "missing_keywords": [str(s) for s in result.get("missing_keywords", []) or []],
        "improvements": [i for i in (result.get("improvements") or []) if isinstance(i, dict)],
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str, model: str) -> dict:
    client = genai.Client(api_key=api_key)
    prompt = f"{SCHEMA_HINT}\n\n"
    if job_description.strip():
        prompt += (
            "Score the resume against this JOB DESCRIPTION. Missing keywords must come "
            f"from it.\n<job_description>\n{job_description[:8000]}\n</job_description>\n\n"
        )
    else:
        prompt += ("No job description was given: judge general ATS-friendliness and "
                   "suggest commonly expected keywords for the candidate's field.\n\n")
    prompt += f"<resume>\n{resume_text[:MAX_CHARS]}\n</resume>"

    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize(parse_json(response.text))


# --------------------------------------------------------------------------- UI
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def get_api_key(user_input: str) -> str:
    if user_input.strip():
        return user_input.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.getenv("GEMINI_API_KEY", "")


def main():
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    with st.sidebar:
        st.header("Settings")
        key_input = st.text_input("Gemini API key", type="password",
                                  help="Get a free key at aistudio.google.com. "
                                       "Not needed if set in secrets.")
        model = st.text_input("Model", value=DEFAULT_MODEL)
        st.markdown("Your resume is sent to Google's Gemini API for analysis. "
                    "Nothing is stored by this app.")

    col1, col2 = st.columns(2)
    with col1:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with col2:
        job_desc = st.text_area("Job description (optional, improves accuracy)", height=150)

    clicked = st.button("Analyze resume", type="primary")
    if not clicked:
        if "result" in st.session_state:
            render_results(st.session_state["result"], st.session_state["checks"])
        return
    if uploaded is None:
        st.warning("Please upload a resume first.")
        return
    api_key = get_api_key(key_input)
    if not api_key:
        st.error("Please provide a Gemini API key in the sidebar.")
        return

    try:
        text = extract_text(uploaded.getvalue(), uploaded.name)
    except Exception as e:
        st.error(f"Could not read the file: {e}")
        return
    if len(text) < 100:
        st.error("Very little text could be extracted. If your PDF is a scanned image, "
                 "ATS systems can't read it either - export a text-based PDF or DOCX.")
        return

    with st.spinner("Analyzing your resume..."):
        try:
            result = analyze_resume(text, job_desc, api_key, model.strip() or DEFAULT_MODEL)
        except Exception as e:
            st.error(f"Analysis failed: {e}")
            return
    # Keep results across reruns (e.g. after clicking the download button)
    st.session_state["result"] = result
    st.session_state["checks"] = basic_checks(text)
    render_results(result, st.session_state["checks"])


def render_results(result: dict, checks: list):
    score = result["overall_score"]
    st.divider()
    left, right = st.columns([1, 2])
    with left:
        st.metric("ATS Score", f"{score}/100", score_label(score), delta_color="off")
        st.progress(score / 100)
    with right:
        st.subheader("Summary")
        st.write(result["summary"])

    st.subheader("Section scores")
    for name, val in result["section_scores"].items():
        st.write(f"**{name.replace('_', ' ').title()}** - {val}/100")
        st.progress(val / 100)

    c1, c2 = st.columns(2)
    with c1:
        st.subheader("✅ Strengths")
        for s in result["strengths"]:
            st.markdown(f"- {s}")
    with c2:
        st.subheader("🔑 Missing keywords")
        if result["missing_keywords"]:
            st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("None found.")

    st.subheader("🛠️ Suggested improvements")
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for imp in result["improvements"]:
        pr = str(imp.get("priority", "Medium"))
        title = f"{icons.get(pr.lower(), '🟠')} {pr} - {imp.get('section', 'General')}"
        with st.expander(title):
            st.markdown(f"**Issue:** {imp.get('issue', '')}")
            st.markdown(f"**Fix:** {imp.get('suggestion', '')}")
            if imp.get("example"):
                st.markdown(f"**Example:** _{imp['example']}_")

    st.subheader("🔍 Quick checks")
    for ok, msg in checks:
        st.write(("✅ " if ok else "❌ ") + msg)

    st.download_button("Download report (JSON)", json.dumps(result, indent=2),
                       file_name="ats_report.json", mime="application/json")


if __name__ == "__main__":
    main()
